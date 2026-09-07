import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

def salvar_com_largura_ajustada(tratado: pd.DataFrame, caminho_arquivo: Path, sheet_name: str) -> None:
    """Salva o DataFrame em Excel ajustando a largura das colunas dinamicamente."""
    caminho_arquivo.parent.mkdir(parents=True, exist_ok=True)
    tratado.to_excel(caminho_arquivo, index=False, sheet_name=sheet_name, engine="openpyxl")

    wb = load_workbook(caminho_arquivo)
    ws = wb[sheet_name]

    for i, coluna in enumerate(tratado.columns, start=1):
        tamanho_maximo = max(
            len(str(coluna)),
            tratado[coluna].map(lambda x: len(str(x))).max(),
        )
        ws.column_dimensions[get_column_letter(i)].width = tamanho_maximo + 2

    wb.save(caminho_arquivo)
    print(f"✅ Dado tratado salvo com sucesso em: {caminho_arquivo}")