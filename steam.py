"""
ETL simples de dados de jogos da Steam.

Fluxo:
1. Requisita os detalhes de um app na Steam API.
2. Salva a resposta bruta (sem nenhuma alteração) em disco -> camada raw.
3. Normaliza o JSON em DataFrame e seleciona só as colunas de interesse.
4. Trata/transforma essas colunas (limpa HTML, extrai requisitos de sistema,
   converte bool, etc) -> camada tratada.
"""
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook
import json
import re
from pathlib import Path

import pandas as pd
import requests
from bs4 import BeautifulSoup

# ---------------------------------------------------------------------------
# Configurações
# ---------------------------------------------------------------------------

APP_ID = 3240220
STEAM_API_URL = f"https://store.steampowered.com/api/appdetails?appids={APP_ID}&cc=br&l=portuguese"

PASTA_DADOS = Path("dados")
CAMINHO_BRUTO = PASTA_DADOS / f"{APP_ID}_bruto.json"

# Colunas que a gente efetivamente quer manter da resposta normalizada.
# Os requisitos de PC/Mac/Linux entram aqui como "matéria-prima": eles são
# usados por extrair_requisitos_steam() e depois descartados em orquestradora().
COLUNAS_DE_INTERESSE = [
    "name", "required_age", "genres",
    "pc_requirements.minimum", "pc_requirements.recommended",
    "mac_requirements.minimum", "mac_requirements.recommended",
    "linux_requirements.minimum", "linux_requirements.recommended",
    "price_overview.currency", "price_overview.initial",
    "price_overview.final", "price_overview.discount_percent",
    "price_overview.initial_formatted", "price_overview.final_formatted",
    "platforms.windows", "platforms.mac", "platforms.linux",
    "metacritic.score", "release_date.date",
]

# ---------------------------------------------------------------------------
# Camada bruta (raw)
# ---------------------------------------------------------------------------


def salvar_camada_bruta(dados_json: dict, caminho_arquivo: Path) -> None:
    """Guarda a resposta original da API em disco, sem nenhuma alteração."""
    caminho_arquivo.parent.mkdir(parents=True, exist_ok=True)
    with open(caminho_arquivo, "w", encoding="utf-8") as f:
        json.dump(dados_json, f, ensure_ascii=False, indent=4)
    print(f"✅ Dado bruto salvo com sucesso em: {caminho_arquivo}")


def buscar_dados_steam(url: str, app_id: int) -> dict:
    """
    Faz a requisição na Steam API e devolve só o dicionário 'data' do jogo,
    já validando se o app_id existe e se a resposta veio com sucesso.
    """
    resposta_api = requests.get(url)
    resposta_api.raise_for_status()  # estoura erro se vier status HTTP ruim (4xx/5xx)

    resposta_json = resposta_api.json()
    salvar_camada_bruta(resposta_json, CAMINHO_BRUTO)

    chave_app = str(app_id)
    bloco_app = resposta_json.get(chave_app)

    if bloco_app is None or not bloco_app.get("success"):
        raise ValueError(f"Steam API não retornou dados válidos para o app_id {app_id}")

    return bloco_app["data"]


# ---------------------------------------------------------------------------
# Funções de tratamento (camada tratada)
# ---------------------------------------------------------------------------


def limpar_html(texto: str) -> str:
    """Remove tags HTML de uma célula de texto, se ela for string."""
    if isinstance(texto, str):
        return BeautifulSoup(texto, "html.parser").get_text(separator=" ").strip()
    return texto


def tratar_bool(infojogo: pd.DataFrame) -> pd.DataFrame:
    """Converte colunas booleanas (True/False) para inteiro (1/0)."""
    colunas_bool = infojogo.select_dtypes(include=["bool"]).columns
    infojogo[colunas_bool] = infojogo[colunas_bool].astype(int)
    return infojogo


def tratar_genero(infojogo: pd.DataFrame) -> str:
    """
    A coluna 'genres' vem como uma lista de dicionários, ex:
    [{"id": "1", "description": "Ação"}, {"id": "25", "description": "Aventura"}]

    Aqui a gente extrai só o campo 'description' de cada dicionário e junta
    tudo numa única string separada por vírgula: "Ação, Aventura".

    Observação: como hoje o script trata um app_id por vez, o DataFrame tem
    sempre uma linha só, por isso acessamos direto o índice 0.
    """
    coluna_genero = infojogo["genres"][0]

    if isinstance(coluna_genero, list):
        generos = [dicionario["description"] for dicionario in coluna_genero]
        return ", ".join(generos)

    return coluna_genero


def extrair_requisitos_steam(infojogo: pd.DataFrame) -> pd.DataFrame:
    """
    Extrai SO, Processador, RAM, GPU e Espaço em disco das colunas de
    requisitos da Steam (que vêm como um único bloco de texto por
    plataforma/nível de exigência).
    """
    # de -> para: nome da coluna original na API -> prefixo da coluna tratada
    colunas_origem = {
        "pc_requirements.minimum": "pc_min",
        "pc_requirements.recommended": "pc_rec",
        "mac_requirements.minimum": "mac_min",
        "mac_requirements.recommended": "mac_rec",
        "linux_requirements.minimum": "linux_min",
        "linux_requirements.recommended": "linux_rec",
    }

    # regex ancorados nos rótulos padrão que a Steam usa nesse texto
    padroes = {
        "so": r"Sistema Operativo:\s*([^:]+?)(?=\s*Processador:|$)",
        "cpu": r"Processador:\s*([^:]+?)(?=\s*Memória:|$)",
        "ram_gb": r"Memória:\s*(\d+)\s*GB",
        "gpu": r"Placa gráfica:\s*([^:]+?)(?=\s*Espaço|$)",
        "disco_gb": r"(?:Espaço no disco:|espaço livre:)?\s*Requer\s*(\d+)\s*GB",
    }

    for col_origem, prefixo in colunas_origem.items():
        if col_origem not in infojogo.columns:
            continue

        infojogo[f"{prefixo}_so"] = (
            infojogo[col_origem].str.extract(padroes["so"], flags=re.IGNORECASE, expand=False).str.strip()
        )
        infojogo[f"{prefixo}_cpu"] = (
            infojogo[col_origem].str.extract(padroes["cpu"], flags=re.IGNORECASE, expand=False).str.strip()
        )

        ram_extraida = infojogo[col_origem].str.extract(padroes["ram_gb"], flags=re.IGNORECASE, expand=False)
        infojogo[f"{prefixo}_ram_gb"] = pd.to_numeric(ram_extraida, errors="coerce")

        infojogo[f"{prefixo}_gpu"] = (
            infojogo[col_origem].str.extract(padroes["gpu"], flags=re.IGNORECASE, expand=False).str.strip()
        )

        disco_extraido = infojogo[col_origem].str.extract(padroes["disco_gb"], flags=re.IGNORECASE, expand=False)
        infojogo[f"{prefixo}_disco_gb"] = pd.to_numeric(disco_extraido, errors="coerce")

            
    return infojogo


def orquestradora(infojogo: pd.DataFrame) -> pd.DataFrame:
    """Executa, em ordem, todas as etapas de tratamento sobre o DataFrame."""

    # 1. Limpa HTML de todas as colunas de texto, exceto 'genres'
    #    (genres tem um tratamento próprio, não é HTML puro)
    colunas_texto = (
        infojogo.select_dtypes(include=["string", "object"])
        .drop(columns=["genres"], errors="ignore")
        .columns
    )
    infojogo[colunas_texto] = infojogo[colunas_texto].map(limpar_html)

    # 2. Transforma a lista de dicionários de 'genres' numa string única
    infojogo["genres"] = tratar_genero(infojogo)

    # 3. Converte colunas booleanas para 0/1
    infojogo = tratar_bool(infojogo)

    # 4. Cria as colunas tratadas de requisitos de sistema (pc_min_so, pc_min_cpu, ...)
    infojogo = extrair_requisitos_steam(infojogo)

    # 5. Remove as colunas antigas de requisitos, já substituídas pelas tratadas
    colunas_antigas = [
        "pc_requirements.minimum", "pc_requirements.recommended",
        "mac_requirements.minimum", "mac_requirements.recommended",
        "linux_requirements.minimum", "linux_requirements.recommended",
    ]
    infojogo = infojogo.drop(columns=colunas_antigas)

    return infojogo


# ---------------------------------------------------------------------------
# Execução
# ---------------------------------------------------------------------------


def main() -> None:
    dados = buscar_dados_steam(STEAM_API_URL, APP_ID)

    tabela_dados = pd.json_normalize(dados)
    infojogo = tabela_dados.reindex(columns=COLUNAS_DE_INTERESSE)

    tratado = orquestradora(infojogo)

    colunas_requisitos = [c for c in tratado.columns if c.startswith(("pc_", "mac_", "linux_"))]
    tratado[colunas_requisitos] = tratado[colunas_requisitos].fillna("Não suportado")


    with pd.option_context("display.max_colwidth", None, "display.max_columns", None):
        for coluna, valor in tratado.iloc[0].items():
            print(f"[{coluna}]:\n{valor}\n{'-' * 40}")
    try:
        tratado.to_excel("xlsx_tratado/dados_jogo.xlsx", index=False, sheet_name=f"jogo ID{APP_ID}", engine="openpyxl")
        print("foi")
    except Exception as e:
        print(e)

    wb = load_workbook("xlsx_tratado/dados_jogo.xlsx")
    ws = wb[f"jogo ID{APP_ID}"]

    for i, coluna in enumerate(tratado.columns, start=1):
        tamanho_maximo = max(
            len(str(coluna)),
            tratado[coluna].map(lambda x: len(str(x))).max())
        ws.column_dimensions[get_column_letter(i)].width = tamanho_maximo + 2

    wb.save("xlsx_tratado/dados_jogo.xlsx")
if __name__ == "__main__":
    main()